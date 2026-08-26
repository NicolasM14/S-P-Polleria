# -*- coding: utf-8 -*-
"""Genera Manual de Usuario profesional S&F Pollería (.docx)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Inches

NAVY = RGBColor(0x00, 0x28, 0x55)
RED = RGBColor(0xE3, 0x1B, 0x23)
GRAY = RGBColor(0x4A, 0x4A, 0x4A)
LIGHT = RGBColor(0x6B, 0x72, 0x80)

OUT = Path(r"c:\Users\nicol\Downloads\Manual_de_Usuario_SF_Polleria.docx")
OUT_PROJECT = Path(r"c:\Users\nicol\OneDrive\Documentos\S&F Polleria\docs\Manual_de_Usuario_SF_Polleria.docx")
LOGO = Path(r"c:\Users\nicol\OneDrive\Documentos\S&F Polleria\public\brand\logo-sf-polleria.jpg")


def set_run_font(run, size=11, bold=False, color=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade_paragraph(paragraph, fill="F3F4F6"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    pPr.append(shd)


def add_horizontal_line(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "002855")
    pBdr.append(bottom)
    pPr.append(pBdr)


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        if level == 1:
            set_run_font(run, 18, True, NAVY)
        elif level == 2:
            set_run_font(run, 14, True, NAVY)
        else:
            set_run_font(run, 12, True, RED)
    return h


def para(doc, text, *, size=11, bold=False, color=GRAY, space_after=8, center=False):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, size, bold, color)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_run_font(run, 11, False, GRAY)


def steps(doc, items):
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Cm(0.3)
        r1 = p.add_run(f"Paso {i}. ")
        set_run_font(r1, 11, True, NAVY)
        r2 = p.add_run(item)
        set_run_font(r2, 11, False, GRAY)


def tip(doc, text, title="Importante"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Cm(0.2)
    p.paragraph_format.right_indent = Cm(0.2)
    shade_paragraph(p, "FEF2F2")
    r1 = p.add_run(f"{title}: ")
    set_run_font(r1, 11, True, RED)
    r2 = p.add_run(text)
    set_run_font(r2, 11, False, GRAY)


def note(doc, text, title="Nota"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    shade_paragraph(p, "EFF6FF")
    r1 = p.add_run(f"{title}: ")
    set_run_font(r1, 11, True, NAVY)
    r2 = p.add_run(text)
    set_run_font(r2, 11, False, GRAY)


def result(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    shade_paragraph(p, "ECFDF5")
    r1 = p.add_run("Resultado: ")
    set_run_font(r1, 11, True, RGBColor(0x04, 0x78, 0x57))
    r2 = p.add_run(text)
    set_run_font(r2, 11, False, GRAY)


def capture_box(doc, number, title, what):
    """Placeholder limpio para insertar captura."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    shade_paragraph(p, "F8FAFC")
    r = p.add_run(f"📷  CAPTURA {number} — {title}")
    set_run_font(r, 10, True, NAVY)

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    shade_paragraph(p2, "F8FAFC")
    r2 = p2.add_run(what)
    set_run_font(r2, 9, False, LIGHT)

    # Espacio visual para pegar la imagen
    for _ in range(3):
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(6)
        shade_paragraph(spacer, "F1F5F9")
        spacer.add_run(" ")


def simple_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True

    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, 10, True, RGBColor(0xFF, 0xFF, 0xFF))
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "002855")
        shading.set(qn("w:val"), "clear")
        hdr[i]._tePr = hdr[i]._tc.get_or_add_tcPr()
        hdr[i]._tc.get_or_add_tcPr().append(shading)

    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = ""
            p = cells[c_idx].paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, 10, False, GRAY)
            if r_idx % 2 == 1:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "F8FAFC")
                shd.set(qn("w:val"), "clear")
                cells[c_idx]._tc.get_or_add_tcPr().append(shd)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def build():
    doc = Document()

    # Márgenes cómodos
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    # ========== PORTADA ==========
    for _ in range(2):
        doc.add_paragraph()

    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(LOGO), width=Cm(6.5))

    para(doc, "MANUAL DE USUARIO", size=28, bold=True, color=NAVY, center=True, space_after=4)
    para(doc, "S&F Pollería", size=22, bold=True, color=RED, center=True, space_after=16)
    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_horizontal_line(line)

    para(
        doc,
        "Guía práctica para usar el sistema día a día",
        size=12,
        color=GRAY,
        center=True,
        space_after=20,
    )
    para(doc, "Versión 0.1.0  ·  Agosto 2026", size=11, color=LIGHT, center=True, space_after=6)
    para(doc, "Para dueños y personal del local", size=11, color=LIGHT, center=True, space_after=6)

    doc.add_page_break()

    # ========== CÓMO LEER ==========
    heading(doc, "Cómo leer este manual", 1)
    para(
        doc,
        "Este manual está pensado para usarse con el sistema abierto al lado. "
        "Cada sección explica una pantalla: para qué sirve, qué hacer y qué pasa después.",
    )
    bullets(
        doc,
        [
            "Los pasos van numerados: seguilos en orden.",
            "Los recuadros rojos son avisos importantes.",
            "Los verdes indican el resultado esperado.",
            "Donde diga “CAPTURA”, más adelante se puede pegar una foto de la pantalla real.",
        ],
    )

    heading(doc, "Tres reglas del sistema", 2)
    simple_table(
        doc,
        ["Tema", "Cómo funciona"],
        [
            ["Cantidades", "Se cargan en GRAMOS (como la balanza). Ejemplo: 1200 = 1,2 kg."],
            ["Precio / costo", "Siempre por KILOGRAMO ($ / kg)."],
            ["Pagos", "Solo Efectivo y Transferencia. El efectivo mueve la caja; la transferencia no."],
        ],
    )
    tip(
        doc,
        "Si la balanza marca 1,250 kg, en el sistema escribí 1250 (gramos). No escribas 1,25.",
    )

    doc.add_page_break()

    # ========== ÍNDICE ==========
    heading(doc, "Contenido", 1)
    toc = [
        "1. Introducción",
        "2. Entrar y salir",
        "3. Inicio (pantalla principal)",
        "4. Menú de navegación",
        "5. Productos",
        "6. Stock",
        "7. Compras",
        "8. Ventas",
        "9. Caja",
        "10. Gastos",
        "11. Reportes",
        "12. Errores frecuentes",
        "13. Buenas prácticas",
        "Anexo — Lista de capturas para completar el PDF",
    ]
    for item in toc:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_run_font(run, 12, False, NAVY)

    doc.add_page_break()

    # ========== 1 ==========
    heading(doc, "1. Introducción", 1)
    heading(doc, "¿Para qué sirve?", 2)
    para(
        doc,
        "S&F Pollería es el panel interno del local. Con él registrás compras, ventas, stock, "
        "caja y gastos, y consultás un resumen del día o del período.",
    )
    heading(doc, "¿Qué podés gestionar?", 2)
    simple_table(
        doc,
        ["Sección", "Para qué"],
        [
            ["Productos", "Catálogo, precio por kg y stock mínimo"],
            ["Stock", "Cuántos gramos hay, ajustes y movimientos"],
            ["Compras", "Ingreso de mercadería y pagos"],
            ["Ventas", "Venta del mostrador y anulación"],
            ["Caja", "Abrir, ver efectivo esperado y cerrar"],
            ["Gastos", "Gastos del local"],
            ["Reportes", "Totales por fechas"],
        ],
    )
    note(
        doc,
        "No es una tienda online ni un sistema de facturas electrónicas. Es gestión interna del local.",
    )

    # ========== 2 ==========
    heading(doc, "2. Entrar y salir", 1)
    heading(doc, "Iniciar sesión", 2)
    steps(
        doc,
        [
            "Abrí el sistema en el navegador (dirección que te den).",
            "Completá Email y Contraseña.",
            "Presioná Ingresar.",
        ],
    )
    result(doc, "Si está bien, aparece “Bienvenido” y entrás a Inicio.")
    tip(
        doc,
        "Si falla, verás: “No se pudo iniciar sesión. Verificá email y contraseña.” "
        "No hay botón de recuperar clave en el sistema: pedila a quien administra los usuarios.",
    )
    heading(doc, "Salir", 2)
    para(doc, "Arriba a la derecha, botón Salir.")
    capture_box(
        doc,
        "01",
        "Login",
        "Pantalla completa de inicio de sesión: logo, Email, Contraseña y botón Ingresar.",
    )

    # ========== 3 ==========
    heading(doc, "3. Inicio (pantalla principal)", 1)
    para(doc, "Es lo primero que ves al entrar. Resume el día.")
    simple_table(
        doc,
        ["Tarjeta", "Qué significa"],
        [
            ["Ventas de hoy", "Total en $ de ventas completadas del día (sin anuladas) y cuántas fueron."],
            ["Caja", "Saldo esperado de efectivo, o “Sin caja abierta”."],
            ["Stock bajo", "Cuántos productos están en el mínimo o por debajo."],
        ],
    )
    para(doc, "También hay accesos rápidos: Nueva venta, Nueva compra, Caja y Productos.")
    tip(doc, "Si dice “Sin caja abierta” y vas a cobrar en efectivo, abrí la caja primero.")
    capture_box(
        doc,
        "02",
        "Inicio",
        "Pantalla Inicio con las 3 tarjetas y los accesos rápidos. Menú lateral visible.",
    )

    # ========== 4 ==========
    heading(doc, "4. Menú de navegación", 1)
    para(doc, "A la izquierda (en computadora) están estas pestañas:")
    simple_table(
        doc,
        ["Pestaña", "Para qué"],
        [
            ["Inicio", "Resumen del día"],
            ["Productos", "Catálogo y precios"],
            ["Stock", "Saldos y ajustes"],
            ["Compras", "Ingreso de mercadería"],
            ["Ventas", "Ventas y anulaciones"],
            ["Caja", "Efectivo del turno"],
            ["Gastos", "Gastos del local"],
            ["Reportes", "Totales por fechas"],
        ],
    )
    tip(
        doc,
        "En pantallas muy chicas (celular) el menú lateral no aparece. Conviene usar una PC o tablet ancha.",
    )
    capture_box(doc, "03", "Menú", "Barra izquierda completa con las 8 opciones y el logo.")

    # ========== 5 ==========
    heading(doc, "5. Productos", 1)
    para(
        doc,
        "Acá está el catálogo. El stock se muestra en gramos. El precio se carga por kilogramo.",
    )

    heading(doc, "Ver y buscar", 2)
    steps(
        doc,
        [
            "Entrá a Productos.",
            "Opcional: buscá por nombre y filtrá por Activos / Inactivos / Todos.",
            "Presioná Filtrar (o Limpiar para sacar filtros).",
        ],
    )
    result(doc, "La tabla muestra Nombre, Precio / kg, Stock (g), Estado y Acciones.")

    heading(doc, "Crear categoría", 2)
    steps(
        doc,
        [
            "En la sección Categorías, escribí el nombre (ej. Elaborados).",
            "Presioná Agregar.",
        ],
    )
    result(doc, "Aparece “Categoría creada”.")

    heading(doc, "Crear producto", 2)
    steps(
        doc,
        [
            "Presioná Nuevo producto.",
            "Completá Nombre, Precio por kg y Stock mínimo (g). La unidad es siempre Gramos (g).",
            "Elegí categoría (opcional) y dejá Producto activo marcado.",
            "Presioná Crear producto.",
        ],
    )
    result(
        doc,
        "El producto queda creado con stock 0 g. El sistema te lleva a Stock para cargar mercadería.",
    )
    tip(
        doc,
        "El stock NO se carga al crear el producto. Después usá Stock (Inventario inicial) o Compras.",
    )

    heading(doc, "Editar / activar / desactivar", 2)
    bullets(
        doc,
        [
            "Editar: en Acciones → Editar → Guardar cambios. El stock se ve pero no se cambia desde ahí.",
            "Desactivar / Activar: en Acciones. Un producto inactivo no se usa en ventas/compras nuevas.",
        ],
    )
    capture_box(doc, "04", "Listado de productos", "Tabla con filtros y al menos 2 productos.")
    capture_box(
        doc,
        "05",
        "Nuevo producto",
        "Formulario con aviso de stock 0 g y botón Crear producto. Ejemplo: Pechuga, $8500/kg, mínimo 2000 g.",
    )

    # ========== 6 ==========
    heading(doc, "6. Stock", 1)
    heading(doc, "Ver saldos", 2)
    para(doc, "En Stock → Saldos actuales ves Producto, Stock (g), Mínimo (g) y Estado.")
    note(
        doc,
        "Ejemplo: si hay 20,5 kg, en pantalla ves 20.500 g. Si vendés 1.200 g, quedan 19.300 g.",
    )

    heading(doc, "Ajustar stock", 2)
    para(doc, "Sirve para inventario inicial, mermas o correcciones (no para una venta normal).")
    steps(
        doc,
        [
            "Elegí el producto.",
            "Cantidad en g: positivo suma (ej. 5000), negativo resta (ej. -500). No puede ser 0.",
            "Elegí el motivo (Inventario inicial, Merma / vencido, etc.). Si es “Otro”, escribí el detalle.",
            "Presioná Registrar ajuste.",
        ],
    )
    result(doc, "Mensaje “Ajuste de stock registrado”. El saldo y el historial se actualizan.")

    heading(doc, "Movimientos", 2)
    para(
        doc,
        "Abajo ves los últimos 50 movimientos: Compra, Venta, Anulación o Ajuste, con cantidad en gramos.",
    )
    capture_box(doc, "07", "Saldos de stock", "Tabla de saldos con columnas en gramos.")
    capture_box(
        doc,
        "08",
        "Ajuste de stock",
        "Formulario con producto, cantidad +10000 y motivo Inventario inicial.",
    )

    # ========== 7 ==========
    heading(doc, "7. Compras", 1)
    para(doc, "Registrá mercadería que entra. Sube el stock y registra el pago.")

    heading(doc, "Nueva compra", 2)
    steps(
        doc,
        [
            "Compras → Nueva compra.",
            "Ítems: producto, Cantidad (g), Costo / kg. Podés agregar más ítems.",
            "Pagos: Efectivo y/o Transferencia. La suma debe ser igual al total. Usá “Completar con total” si hay un solo pago.",
            "Notas opcionales (proveedor, remito).",
            "Registrar compra.",
        ],
    )
    result(
        doc,
        "“Compra registrada”. El stock sube. Si pagaste en efectivo y hay caja abierta, baja el efectivo de la caja.",
    )
    tip(
        doc,
        "Costo / kg es lo que te cobró el proveedor, NO el precio de venta al público.",
    )
    note(
        doc,
        "Ejemplo: 10 kg a $6.000/kg → cantidad 10000 g, costo 6000, total $60.000. Los pagos deben sumar $60.000.",
    )
    capture_box(
        doc,
        "11",
        "Nueva compra",
        "Pantalla en dos columnas (Ítems | Pagos) con total y “Pagos = total”.",
    )

    # ========== 8 ==========
    heading(doc, "8. Ventas", 1)
    para(doc, "Venta del mostrador: descuenta stock y registra el cobro.")

    heading(doc, "Nueva venta", 2)
    steps(
        doc,
        [
            "Ventas → Nueva venta (o desde Inicio).",
            "Elegí producto. Se completa el precio/kg y ves el stock en g.",
            "Cantidad (g): lo de la balanza (ej. 1200). Revisá cuánto stock quedaría.",
            "Descuento opcional en pesos. Revisá Subtotal y Total.",
            "Pagos: Efectivo y/o Transferencia hasta igualar el total.",
            "Registrar venta.",
        ],
    )
    result(
        doc,
        "“Venta registrada”. Baja el stock. El efectivo suma a la caja solo si hay caja abierta.",
    )
    note(
        doc,
        "Ejemplo: pechuga a $8.000/kg, vendés 1200 g → 1,2 × 8000 = $9.600. Ese es el total a cobrar (salvo descuento).",
    )

    heading(doc, "Anular una venta", 2)
    steps(
        doc,
        [
            "En el listado de Ventas, en una venta Completada, presioná Anular.",
            "Confirmá el mensaje.",
        ],
    )
    result(
        doc,
        "Pasa a Anulada. Vuelve el stock. Se revierte el efectivo en caja si correspondía.",
    )
    tip(doc, "No se puede editar una venta: si te equivocás, anulá y cargá de nuevo.")
    capture_box(
        doc,
        "14",
        "Nueva venta",
        "Dos columnas con producto, 1200 g, stock disponible y pagos iguales al total.",
    )

    # ========== 9 ==========
    heading(doc, "9. Caja", 1)
    para(doc, "Controla el efectivo físico del turno. Solo una caja abierta a la vez.")

    heading(doc, "Abrir", 2)
    steps(
        doc,
        [
            "Entrá a Caja.",
            "Ingresá el Monto de apertura (lo que hay en el cajón; puede ser 0).",
            "Abrir caja.",
        ],
    )

    heading(doc, "Durante el día", 2)
    simple_table(
        doc,
        ["Qué pasa", "Efecto en caja"],
        [
            ["Venta en efectivo", "Suma"],
            ["Compra en efectivo", "Resta"],
            ["Gasto “Descontar de caja”", "Resta"],
            ["Todo lo que es transferencia", "No mueve la caja"],
        ],
    )
    para(doc, "Vas a ver el Saldo esperado: lo que el sistema calcula que debería haber en efectivo.")

    heading(doc, "Cerrar", 2)
    steps(
        doc,
        [
            "Contá el efectivo del cajón.",
            "Ingresá Efectivo contado.",
            "Revisá la diferencia respecto del esperado.",
            "Notas opcionales y confirmá el cierre.",
        ],
    )
    result(doc, "“Caja cerrada”. Queda en Sesiones cerradas recientes.")
    capture_box(doc, "18", "Caja abierta", "Saldo esperado y tabla de movimientos del turno.")

    # ========== 10 ==========
    heading(doc, "10. Gastos", 1)
    steps(
        doc,
        [
            "Gastos → Nuevo gasto.",
            "Elegí categoría, monto, descripción y fecha.",
            "Si el dinero salió del cajón, marcá “Descontar de caja (efectivo)”.",
            "Registrar gasto.",
        ],
    )
    result(doc, "“Gasto registrado”. Si marcaste caja y hay sesión abierta, baja el efectivo.")
    capture_box(doc, "22", "Nuevo gasto", "Formulario con categoría, monto y checkbox de caja.")

    # ========== 11 ==========
    heading(doc, "11. Reportes", 1)
    steps(
        doc,
        [
            "Entrá a Reportes.",
            "Elegí Desde y Hasta, o usá Hoy.",
            "Presioná Filtrar.",
        ],
    )
    para(
        doc,
        "Vas a ver totales de Ventas (sin anuladas), Compras, Gastos, estado de Caja y productos con stock bajo.",
    )
    capture_box(doc, "23", "Reportes", "Totales del período y, si hay, tabla de stock bajo.")

    # ========== 12 ==========
    heading(doc, "12. Errores frecuentes", 1)
    simple_table(
        doc,
        ["Mensaje / situación", "Qué hacer"],
        [
            ["No se pudo iniciar sesión…", "Revisá email y contraseña."],
            ["Revisá las cantidades en gramos…", "Usá enteros: 1200, no 1,2."],
            ["Stock insuficiente…", "Comprá o ajustá stock antes de vender."],
            ["La suma de los pagos debe coincidir…", "Igualá pagos al total (Completar con total)."],
            ["Ya hay una caja abierta", "Usá la actual o cerrala primero."],
            ["No hay caja abierta", "Abrí caja en la pestaña Caja."],
            ["Esa venta ya fue anulada", "No hace falta anular otra vez."],
            ["Error raro / JWT / hora", "Sincronizá la hora de Windows y volvé a entrar."],
        ],
    )

    # ========== 13 ==========
    heading(doc, "13. Buenas prácticas", 1)
    bullets(
        doc,
        [
            "Abrí la caja al empezar el turno si vas a manejar efectivo.",
            "Cargá el stock inicial en Stock o Compras, no en el alta del producto.",
            "Escribí siempre gramos de la balanza.",
            "Antes de guardar, mirá que diga “Pagos = total”.",
            "Revisá stock bajo en Inicio o Reportes.",
            "Si te equivocás en una venta: Anular y cargar de nuevo.",
            "Al cerrar, contá el efectivo y compará con el esperado.",
            "Usá el sistema en computadora para ver el menú completo.",
        ],
    )

    # ========== ANEXO ==========
    doc.add_page_break()
    heading(doc, "Anexo — Capturas para completar el documento", 1)
    para(
        doc,
        "Esta lista es para quien arma el PDF final. Pegá cada captura en el lugar marcado "
        "con el mismo número.",
    )
    simple_table(
        doc,
        ["N.º", "Pantalla", "Qué mostrar"],
        [
            ["01", "Login", "Email, contraseña, Ingresar"],
            ["02", "Inicio", "3 tarjetas + accesos"],
            ["03", "Menú", "8 pestañas"],
            ["04", "Productos", "Listado con stock en g"],
            ["05", "Nuevo producto", "Formulario completo"],
            ["07", "Stock", "Saldos"],
            ["08", "Stock", "Ajuste con motivo"],
            ["11", "Nueva compra", "Ítems + pagos"],
            ["14", "Nueva venta", "1200 g + pagos"],
            ["18", "Caja", "Sesión abierta"],
            ["22", "Nuevo gasto", "Formulario"],
            ["23", "Reportes", "Totales del período"],
        ],
    )
    tip(
        doc,
        "Si usás el sistema real del local, no cambies datos solo para sacar fotos. "
        "Preferí un entorno de prueba o capturas con datos ya existentes.",
        title="Aviso",
    )

    para(
        doc,
        "— Fin del manual — S&F Pollería —",
        size=10,
        color=LIGHT,
        center=True,
        space_after=0,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT_PROJECT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    doc.save(OUT_PROJECT)
    print(f"OK: {OUT}")
    print(f"OK: {OUT_PROJECT}")


if __name__ == "__main__":
    build()
